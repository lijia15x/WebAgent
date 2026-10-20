import asyncio
import json
import re
import unicodedata
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

from common.skills.copilot_sdk.copilot_sdk_client import ContextTier, ReasoningEffort, ask_copilot

from .artifact_store import WORKSPACE_ROOT, register_ppt
from .config import MrcConfig
from .models import MrcArtifact


SUMMARY_FONT_SIZE_PT = 10
PAGE3_FONT_SIZE_PT = 10
PAGE3_MAX_CHAR_COUNT = 4900
PAGE2_DETAIL_TEXTBOX_LEFT = 588135
PAGE2_DETAIL_TEXTBOX_TOP = 2871854
PAGE2_DETAIL_TEXTBOX_WIDTH = 11168038
PAGE2_DETAIL_TEXTBOX_HEIGHT = 2808461
DEFAULT_TEMPLATE_PATH = (
    Path(__file__).resolve().parent
    / "templates"
    / "DMR Johnson City SXT Executive Summary V2.pptx"
)
DMR_RS_TEMPLATE_PATH = (
    Path(__file__).resolve().parent / "templates" / "DMR-RS_template.pptx"
)
DMR_16CH_TEMPLATE_PATH = (
    Path(__file__).resolve().parent / "templates" / "DMR-16ch_template.pptx"
)
TEMPLATE_PATH = DEFAULT_TEMPLATE_PATH

DASHBOARD_PAGE2_PROMPT = """Prepare the DMR-16ch one-page MRC executive pre-read for {cycle_code}.
It will be sent one day before a 30-minute project review. Summarize facts, not generic leadership judgments.
Workbook worksheets: {sheet_names}

Identify the execution-status worksheet using its actual headers and content. A stale worksheet name or
project label is not proof that all rows are stale. Use explicit current-week headings or dated narrative
where available. If the workbook has a single undated Status Comments column, treat it as the supplied
snapshot, do not claim every row was updated this week, and do not infer week-over-week changes.
Use the actual worksheet name in sources, even when its week differs from the reporting cycle.
Carry merged or blank Platform cells downward. Platform is the engineering group;
Engineering Domain is the project or domain within that group. Identify the Domain in each update.

Read the workbook's status legend. Interpret independent dimensions separately:
R = Blocked (no mitigation plan); O = At Risk (mitigation in progress);
Y = On Track (approved mitigation); G = On Track (no issue); Done = completed; ignore N/A.
Do not call Y blocked or at risk. Missing status cells are unknown, not green.
Use status and mitigation narrative as evidence. A blank mitigation column does not prove no mitigation
exists if the status narrative documents one. Never invent solutions, owners, deadlines or overall health.
Do not import unrelated COR, other-platform or shared-program facts as DMR-16ch-specific facts;
retain their documented scope and mention them only when a DMR dependency is explicit.

Return only a valid JSON object with these keys and schema:
{{"highlights":[{{"bold_lead":"achievement or material change","normal_detail":"evidence","source":"exact worksheet!L10"}}],
"group_updates":[{{"title":"Engineering group","paragraphs":[
{{"bold_lead":"Domain and issue","normal_detail":"impact; documented mitigation or pending support","source":"exact worksheet!L10"}}
]}}],
"milestones":[{{"name":"deliverable or event","drive_to":null,"drive_to_source":"",
"trend_actual":null,"trend_kind":null,"trend_source":"","source":"exact worksheet!L10"}}]}}

Requirements:
- Select up to 2 Highlights, up to 4 engineering groups and up to 4 group-update bullets IN TOTAL.
- Do not force a fixed number of groups or bullets. Prioritize material issues, then consolidate healthy progress.
- Each bullet's bold_lead plus normal_detail must be at most 180 characters including the separating space.
- Each group title must be at most 60 characters. Use concise English; no Markdown, URLs or visible cell addresses.
- Include up to 5 evidenced milestones; use an empty list when none are documented.
- Milestones may be delivery or qualification events from narrative, not only formal project phase gates.
- name: at most 48 characters. drive_to: explicit committed/requested TARGET date, at most 24 characters, or null.
- trend_actual: explicit forecast, actual or worst-case date, at most 40 characters, or null.
- trend_kind: "forecast", "actual", "worst_case", or null when trend_actual is null.
- Date fields contain only date text; put forecast/actual/worst-case labels exclusively in trend_kind.
- Preserve dates verbatim as substrings of their cited cells. Do not infer a year or convert ranges to exact weeks.
- drive_to_source and trend_source must cite the individual nonblank cells containing those date strings.
- A worst-case date must remain worst_case, not a commitment. Do not infer dates from completion percentages.
- Every bullet and milestone must cite an existing nonblank cell, using the exact worksheet name and one A1 address.
- Never cite blank cells, merged placeholders, or historical text as a current-week update.
- Missing dates must be null, not guessed or copied from the PPT template. No unsupported milestone colors.
- Choose explicit support requests where present; otherwise describe pending facts without inventing a decision request.
- Escape quotes, backslashes and newlines inside JSON strings. No commentary or code fences."""

DASHBOARD_PAGE3_PROMPT = """Prepare concise supporting MRC detail for DMR-16ch, reporting cycle {cycle_code}.
Workbook worksheets: {sheet_names}
Organize the narrative by Platform (engineering group), identifying Engineering Domains in the facts.
Carry merged Platform cells downward. Use the workbook's current reporting columns when present.
If only one undated Status Comments column exists, treat it as the supplied snapshot, not verified weekly change.
Worksheet names may be stale: cite the exact actual name and do not rename source sheets to match the cycle.
Use current evidence, not historical text as a current update. Retain the scope of other-project dependencies.
R = Blocked, O = At Risk, Y = On Track with approved mitigation, G = On Track, Done = completed; ignore N/A.
Missing structured status is unknown, not green. Do not infer overall project health.
Prioritize unresolved issues, impact, documented measures and pending support; keep healthy progress brief.
Never invent solutions, owners, dates, status changes or decisions. Keep date ranges and worst-case qualifiers.
Return only valid JSON:
{{"sections":[{{"title":"Engineering group","paragraphs":[
{{"bold_lead":"Domain / fact","normal_detail":"supporting facts and documented actions","source":"exact worksheet!L10"}}
]}}]}}
Use up to 5 groups and up to 6 bullets IN TOTAL, with at most 260 characters per bullet (lead plus detail).
Each title is at most 60 characters. Do not pad groups to two paragraphs or force coverage of empty rows.
Every bullet needs a real nonblank source cell using the exact worksheet name and one A1 address.
Keep narrative within {max_chars} characters. Use concise English, no Markdown, URLs or code fences."""

PAGE2_PROMPT = """You are preparing slide 2 of an executive summary PPT from the attached Excel workbook.

Reporting cycle: {cycle_code}
Workbook worksheets: {sheet_names}

Identify the execution-status worksheet and the current reporting-week columns. Convert {cycle_code} from YYYYWWNN
to the workbook label WWNN'YY when necessary. For example, 2026WW40 corresponds to WW40'26.
A merged week heading may cover two columns:
- Status Comments / Status & Problem Statement
- Mitigation Plan

Use the Platform and Engineering Domain columns as the reporting hierarchy. Platform is the engineering group;
Engineering Domain is the project or domain within that group. Carry merged or blank Platform cells downward until
the next non-empty Platform value.

Treat the workbook's color-grading key as authoritative:
- R = Blocked: critical issue with no mitigation plan
- O = At Risk: mitigation plan is still work in progress
- Y = On Track: approved mitigation plan exists
- G = On Track: no issue
- Done = all deliverables completed
- N/A = not applicable

The status columns such as Schedule, Resource and Spending, Platform PCOS, NUDDs, Quality, and Technical Execution
are independent health dimensions. Rank attention using R > O > Y > G > Done; ignore N/A.
Do not describe Y as blocked or at risk.

Use the current-week status and mitigation columns as the primary narrative evidence. Use older weekly columns only
to identify an explicit trend or change. If the current-week status cell is blank, state that no current-week
narrative update was provided; never substitute historical text as a current update.

Return only valid JSON with this schema:
{{"executive_summary":[
  {{"bold_lead":"leadership judgment","normal_detail":"supporting workbook facts"}},
  {{"bold_lead":"leadership judgment","normal_detail":"supporting workbook facts"}}
],
"detail_sections":[
  {{"title":"Platform name","paragraphs":[
    {{"bold_lead":"priority judgment","normal_detail":"supporting Domain, status dimension, current-week facts and mitigation"}},
    {{"bold_lead":"priority judgment","normal_detail":"supporting Domain, status dimension, current-week facts and mitigation"}}
  ]}}
]}}

Requirements:
- Produce exactly 2 executive_summary paragraphs.
- Select exactly 2 Platforms requiring the most leadership attention.
- Produce exactly 2 paragraphs for each selected Platform.
- Name the relevant Engineering Domain and status dimension in supporting details.
- Prioritize blocked and at-risk items, missing mitigation, schedule impact, dependencies, and dated milestones.
- Consolidate healthy progress instead of listing every green item.
- Do not invent facts, dates, quantities, status meanings, risks, causes, or mitigations.
- Do not mention cell addresses in slide 2.
- Do not use Markdown or code fences."""

PAGE3_PROMPT = """Create slide 3 detailed executive narrative for {cycle_code} from the attached Excel workbook.

Workbook worksheets: {sheet_names}

Identify the execution-status worksheet and the current reporting-week columns. Convert {cycle_code} from YYYYWWNN
to the workbook label WWNN'YY when necessary. For example, 2026WW40 corresponds to WW40'26.
A merged week heading may cover the Status & Problem Statement column and its Mitigation Plan column.

Organize the narrative by Platform and describe the Engineering Domain rows within each Platform.
Carry merged or blank Platform cells downward until the next non-empty Platform value.

Treat the workbook's color-grading key as authoritative:
- R = Blocked: critical issue with no mitigation plan
- O = At Risk: mitigation plan is still work in progress
- Y = On Track: approved mitigation plan exists
- G = On Track: no issue
- Done = all deliverables completed
- N/A = not applicable

Evaluate Schedule, Resource and Spending, Platform PCOS, NUDDs, Quality, and Technical Execution independently.
Use severity R > O > Y > G > Done and ignore N/A. Do not describe Y as blocked or at risk.

Use the current-week Status & Problem Statement and Mitigation Plan cells as primary evidence.
Historical week columns may be used only for an explicit trend comparison.
If the current-week narrative is blank, say that no current-week update was provided and rely only on the structured
status cells; never present historical text as a current-week update.

Return only valid JSON with this schema:
{{"sections":[
  {{"title":"Platform name","paragraphs":[
    {{"bold_lead":"synthesized conclusion","normal_detail":"supporting Domain, status dimension, current-week facts and mitigation","source":"exact worksheet name!L12"}},
    {{"bold_lead":"synthesized conclusion","normal_detail":"supporting Domain, status dimension, current-week facts and mitigation","source":"exact worksheet name!E12"}}
  ]}}
]}}

Requirements:
- Include every major Platform represented in the execution-status table.
- Produce exactly 2 paragraphs per Platform.
- Prioritize R and O items, then Y items needing monitoring; consolidate G progress.
- Identify the Engineering Domain and affected status dimension.
- Mention concrete dates, quantities, dependencies, milestones, and mitigation only when present in the workbook.
- Keep all returned narrative text within {max_chars} characters.
- Every source must be an existing cell using the exact worksheet name and address.
- Prefer the current-week narrative or mitigation cell as the source.
- When no current-week narrative exists, cite the relevant structured status cell instead.
- Never cite a blank cell, merged placeholder cell, historical cell presented as current, or invented address.
- Do not include URLs, Markdown, or code fences."""


def _parse_json_response(content: str) -> dict[str, Any]:
    cleaned = content.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE)
    payload = json.loads(cleaned)
    if not isinstance(payload, dict):
        raise ValueError("Copilot PPT response must be a JSON object")
    return payload


def _call_copilot_json(
    prompt: str,
    model: str,
    excel_path: Path,
    *,
    reasoning_effort: ReasoningEffort = "high",
    context_tier: ContextTier = "default",
) -> dict[str, Any]:
    content = asyncio.run(
        ask_copilot(
            prompt, model=model, files=[excel_path], workspace="",
            reasoning_effort=reasoning_effort, context_tier=context_tier,
        )
    )
    return _parse_json_response(content)


def _sheet_names(excel_path: Path) -> list[str]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def _template_path_for(file_name: str) -> Path:
    if "dmr-rs" in file_name.casefold():
        return DMR_RS_TEMPLATE_PATH
    if "dmr-16ch" in file_name.casefold():
        return DMR_16CH_TEMPLATE_PATH
    return DEFAULT_TEMPLATE_PATH


def _paragraphs(value: object, count: int | None = None) -> list[dict[str, str]]:
    normalized = []
    if isinstance(value, list):
        for item in value:
            if isinstance(item, dict):
                lead = str(item.get("bold_lead") or item.get("lead") or "").strip()
                detail = str(item.get("normal_detail") or item.get("detail") or "").strip()
                source = str(item.get("source") or item.get("reference") or "").strip()
            else:
                lead, detail, source = str(item).strip(), "", ""
            if lead or detail:
                normalized.append(
                    {"bold_lead": lead, "normal_detail": detail, "source": source}
                )
    if count is not None:
        normalized = normalized[:count]
        while len(normalized) < count:
            normalized.append(
                {
                    "bold_lead": "No update available.",
                    "normal_detail": "",
                    "source": "",
                }
            )
    return normalized


def _prepare_text_frame(shape):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Pt(0)
    return frame


def _source_link(source_url: str, source: str) -> str:
    encoded_source = quote(source, safe="")
    if "?" in source_url:
        return f"{source_url}&activeCell={encoded_source}&wdActiveCell={encoded_source}"
    return f"{source_url}?web=1&activeCell={encoded_source}&wdActiveCell={encoded_source}"


def _add_paragraph(
    frame,
    index: int,
    lead: str,
    detail: str = "",
    level: int = 0,
    source: str = "",
    source_url: str = "",
    source_label: str = "",
) -> int:
    paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
    paragraph.level = level
    paragraph.space_before = Pt(0 if index == 0 else 6)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    for text, bold in ((lead, True), ((" " if lead and detail else "") + detail, False)):
        if text:
            run = paragraph.add_run()
            run.text = text
            run.font.size = Pt(SUMMARY_FONT_SIZE_PT)
            run.font.bold = bold
    if source and source_url:
        run = paragraph.add_run()
        run.text = f" {source_label}" if source_label else f" [{source}]"
        run.font.size = Pt(SUMMARY_FONT_SIZE_PT)
        run.font.underline = True
        if source_label:
            run.font.color.rgb = RGBColor.from_string("24758E")
        run.hyperlink.address = _source_link(source_url, source)
    return index + 1


def _set_sections(
    shape,
    sections: list[dict[str, Any]],
    font_size: int = SUMMARY_FONT_SIZE_PT,
    source_url: str = "",
) -> None:
    frame = _prepare_text_frame(shape)
    index = 0
    for section in sections:
        title = str(section.get("title") or "Engineering Update").strip()
        index = _add_paragraph(frame, index, title)
        for paragraph in _paragraphs(section.get("paragraphs"), count=2):
            index = _add_paragraph(
                frame,
                index,
                paragraph["bold_lead"],
                paragraph["normal_detail"],
                source=paragraph["source"],
                source_url=source_url,
            )
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def _text_shapes(slide) -> list[Any]:
    return [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]


def _first_font_size(shape):
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                return run.font.size
    return None


def _set_text_preserving_font_size(shape, text: str) -> None:
    font_size = _first_font_size(shape)
    shape.text = text
    if font_size is not None:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = font_size


def _replace_week_tokens(slide, label: str) -> None:
    for shape in _text_shapes(slide):
        if shape.text and re.search(r"(?i)WW\s*\d{1,2}", shape.text):
            _set_text_preserving_font_size(
                shape,
                re.sub(
                    r"(?i)WW\s*\d{1,2}(?:\s*[‘’']?\s*\d{2,4})?",
                    label,
                    shape.text,
                ),
            )


def _set_dashboard_sections(
    shape, sections: list[dict[str, Any]], source_url: str, font_size: int = 12
) -> None:
    frame = _prepare_text_frame(shape)
    frame.auto_size = MSO_AUTO_SIZE.NONE
    index = 0
    for section in sections:
        if section.get("title"):
            index = _add_paragraph(frame, index, str(section["title"]))
        for paragraph in _paragraphs(section.get("paragraphs")):
            index = _add_paragraph(
                frame, index, paragraph["bold_lead"], paragraph["normal_detail"],
                source=paragraph["source"], source_url=source_url, source_label="Details",
            )
    if index == 0:
        _add_paragraph(frame, 0, "No current-week update provided.")
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def _update_dashboard(slide, cycle_code: str, page2: dict[str, Any]) -> None:
    shapes = {shape.name: shape for shape in slide.shapes}
    required = {"dashboard_title", "highlights", "group_updates", "milestone_table", "source_workbook_link"}
    if not required <= shapes.keys():
        raise ValueError("Dashboard template is missing named content regions")
    _replace_week_tokens_for_shape(
        shapes["dashboard_title"], f"WW{cycle_code[-2:]}'{cycle_code[:4]}"
    )
    source_url = str(page2.get("source_url") or "")
    _set_dashboard_sections(
        shapes["highlights"], [{"paragraphs": page2.get("highlights", [])}], source_url
    )
    _set_dashboard_sections(shapes["group_updates"], page2.get("group_updates", []), source_url)
    table = shapes["milestone_table"].table
    milestones = page2.get("milestones", [])
    for index, row in enumerate(list(table.rows)[1:]):
        milestone = milestones[index] if index < len(milestones) else {}
        trend = str(milestone.get("trend_actual") or "Not provided")
        trend_label = {"forecast": "Forecast", "actual": "Actual", "worst_case": "Worst case"}.get(
            milestone.get("trend_kind")
        )
        if milestone.get("trend_actual") and trend_label:
            trend = f"{trend_label}: {trend}"
        values = (
            str(milestone.get("name") or ("No sourced milestones" if index == 0 else "")),
            str(milestone.get("drive_to") or ("Not provided" if milestone else "")),
            trend if milestone else "",
        )
        for cell, text in zip(row.cells, values):
            _set_text_preserving_font_size(cell, text)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor.from_string("243238")
        if milestone.get("source") and source_url:
            run = row.cells[0].text_frame.paragraphs[0].add_run()
            run.text = " Details"
            run.font.size = Pt(10)
            run.font.underline = True
            run.font.color.rgb = RGBColor.from_string("24758E")
            run.hyperlink.address = _source_link(source_url, milestone["source"])
        for column, source_key in ((1, "drive_to_source"), (2, "trend_source")):
            if milestone.get(source_key) and source_url:
                for paragraph in row.cells[column].text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.hyperlink.address = _source_link(source_url, milestone[source_key])
                        run.font.underline = True
                        run.font.color.rgb = RGBColor.from_string("24758E")
    workbook_link = shapes["source_workbook_link"]
    _set_text_preserving_font_size(
        workbook_link, "Source Excel" if source_url else "Source Excel unavailable"
    )
    if source_url:
        run = workbook_link.text_frame.paragraphs[0].runs[0]
        run.hyperlink.address = source_url
        run.font.underline = True
        run.font.color.rgb = RGBColor.from_string("24758E")


def _update_page2(slide, cycle_code: str, page2: dict[str, Any]) -> None:
    if any(shape.name == "highlights" for shape in slide.shapes):
        _update_dashboard(slide, cycle_code, page2)
        return
    long_text_shapes = []
    dashboard_label = f"WW{cycle_code[-2:]}'{cycle_code[:4]}"
    for shape in _text_shapes(slide):
        text = (shape.text or "").strip()
        if not text:
            continue
        if re.search(r"(?i)platform\s+dashboard", text):
            _replace_week_tokens_for_shape(shape, dashboard_label)
        elif len(text) > 80:
            long_text_shapes.append(shape)

    if not long_text_shapes:
        raise ValueError("PPT template slide 2 has no summary text area")
    summary = _paragraphs(page2.get("executive_summary"), count=2)
    _set_sections(long_text_shapes[0], [{"title": "Executive Summary", "paragraphs": summary}])

    details = page2.get("detail_sections")
    if isinstance(details, list):
        detail_shape = (
            long_text_shapes[1]
            if len(long_text_shapes) > 1
            else slide.shapes.add_textbox(
                PAGE2_DETAIL_TEXTBOX_LEFT,
                PAGE2_DETAIL_TEXTBOX_TOP,
                PAGE2_DETAIL_TEXTBOX_WIDTH,
                PAGE2_DETAIL_TEXTBOX_HEIGHT,
            )
        )
        _set_sections(detail_shape, details[:2])


def _replace_week_tokens_for_shape(shape, label: str) -> None:
    _set_text_preserving_font_size(
        shape,
        re.sub(
            r"(?i)WW\s*\d{1,2}(?:\s*[‘’']?\s*\d{2,4})?",
            label,
            shape.text,
        ),
    )


def _update_presentation(
    template_path: Path,
    output_path: Path,
    cycle_code: str,
    page2: dict[str, Any],
    page3: dict[str, Any],
) -> None:
    presentation = Presentation(str(template_path))
    while len(presentation.slides) > 3:
        slide_id = list(presentation.slides._sldIdLst)[-1]
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    if len(presentation.slides) < 3:
        raise ValueError("PPT template must contain at least three slides")

    dashboard_label = f"WW{cycle_code[-2:]}'{cycle_code[:4]}"
    _replace_week_tokens(presentation.slides[0], dashboard_label)
    _update_page2(presentation.slides[1], cycle_code, page2)
    _replace_week_tokens(presentation.slides[2], dashboard_label)

    page3_shapes = sorted(
        _text_shapes(presentation.slides[2]), key=lambda shape: len(shape.text or ""), reverse=True
    )
    if not page3_shapes:
        raise ValueError("PPT template slide 3 has no text area")
    sections = page3.get("sections")
    details_shape = next(
        (shape for shape in page3_shapes if shape.name == "details_body"), None
    )
    if details_shape is not None:
        _set_dashboard_sections(
            details_shape, sections if isinstance(sections, list) else [],
            str(page3.get("source_url") or ""), font_size=12,
        )
    else:
        _set_sections(
            page3_shapes[0],
            sections if isinstance(sections, list) else [],
            PAGE3_FONT_SIZE_PT,
            source_url=str(page3.get("source_url") or ""),
        )
    presentation.save(str(output_path))


def _date_in_source(date: str, source: str) -> bool:
    def normalize(value: str) -> str:
        value = unicodedata.normalize("NFKC", value).casefold()
        value = value.translate(str.maketrans({
            "\u2018": "'", "\u2019": "'", "\u02bc": "'",
            "\u2010": "-", "\u2011": "-", "\u2013": "-", "\u2014": "-",
        }))
        value = re.sub(r"\s+", " ", value).strip()
        value = re.sub(r"\s*(['/-])\s*", r"\1", value)
        return re.sub(r"\bww\s+(?=\d)", "ww", value)

    normalized_date = normalize(date)
    return bool(normalized_date) and re.search(
        r"(?<!\w)" + re.escape(normalized_date) + r"(?!\w)", normalize(source)
    ) is not None


def _validate_dashboard_payloads(excel_path: Path, page2: dict[str, Any], page3: dict[str, Any]) -> None:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        def items(value: object, limit: int, label: str) -> list[dict[str, Any]]:
            if not isinstance(value, list) or len(value) > limit:
                raise ValueError(f"{label} must be a list with at most {limit} items")
            if any(not isinstance(item, dict) for item in value):
                raise ValueError(f"{label} items must be JSON objects")
            return value

        def source_text(source: object) -> str:
            if not isinstance(source, str) or "!" not in source:
                raise ValueError("Dashboard evidence requires a worksheet!A1 source")
            sheet_name, address = source.rsplit("!", 1)
            if sheet_name not in workbook.sheetnames or not re.fullmatch(r"[A-Za-z]{1,3}[1-9]\d*", address):
                raise ValueError(f"Invalid dashboard source: {source}")
            sheet = workbook[sheet_name]
            try:
                cell = sheet[address]
            except (ValueError, IndexError) as error:
                raise ValueError(f"Invalid dashboard source: {source}") from error
            if cell.value is None or not str(cell.value).strip():
                raise ValueError(f"Dashboard source is blank: {source}")
            return str(cell.value)

        def bullets(value: object, limit: int, max_chars: int, label: str) -> int:
            paragraphs = items(value, limit, label)
            for paragraph in paragraphs:
                lead, detail = paragraph.get("bold_lead", ""), paragraph.get("normal_detail", "")
                if not isinstance(lead, str) or not isinstance(detail, str):
                    raise ValueError(f"{label} text must be strings")
                text = f"{lead} {detail}".strip()
                if not text or len(text) > max_chars:
                    raise ValueError(f"{label} bullet must contain 1-{max_chars} characters")
                source_text(paragraph.get("source"))
            return len(paragraphs)

        def sections(value: object, max_groups: int, max_bullets: int, max_chars: int, label: str) -> None:
            total = 0
            for section in items(value, max_groups, label):
                title = section.get("title")
                if not isinstance(title, str) or not 1 <= len(title.strip()) <= 60:
                    raise ValueError(f"{label} title must contain 1-60 characters")
                total += bullets(section.get("paragraphs"), max_bullets, max_chars, label)
            if total > max_bullets:
                raise ValueError(f"{label} exceeds {max_bullets} total bullets")

        bullets(page2.get("highlights"), 2, 180, "Highlights")
        sections(page2.get("group_updates"), 4, 4, 180, "Group updates")
        for milestone in items(page2.get("milestones"), 5, "Milestones"):
            name = milestone.get("name")
            if not isinstance(name, str) or not 1 <= len(name.strip()) <= 48:
                raise ValueError("Milestone name must contain 1-48 characters")
            source_text(milestone.get("source"))
            for key, source_key, limit in (("drive_to", "drive_to_source", 24), ("trend_actual", "trend_source", 40)):
                date = milestone.get(key)
                if date is not None:
                    if not isinstance(date, str) or not 1 <= len(date.strip()) <= limit:
                        raise ValueError(f"Milestone {key} must be a date string or null")
                    if not _date_in_source(date, source_text(milestone.get(source_key))):
                        raise ValueError(
                            f"Milestone {name!r} {key} date {date!r} not found in "
                            f"source {milestone.get(source_key)!r}"
                        )
            kind = milestone.get("trend_kind")
            if milestone.get("trend_actual") is not None:
                if not isinstance(kind, str) or kind not in {"forecast", "actual", "worst_case"}:
                    raise ValueError("Milestone trend_kind must identify forecast, actual or worst_case")
            elif kind is not None:
                raise ValueError("Milestone trend_kind must be null when the date is missing")
        sections(page3.get("sections"), 5, 6, 260, "Supporting detail")
    except ValueError as error:
        raise ValueError(f"Invalid DMR-16ch report for {excel_path.name}: {error}") from error
    finally:
        workbook.close()


def _call_dashboard_page(
    prompt: str, config: MrcConfig, excel_path: Path, page_number: int
) -> dict[str, Any]:
    request_prompt = prompt
    max_attempts = 3
    for attempt in range(1, max_attempts + 1):
        payload = _call_copilot_json(
            request_prompt, config.ppt_model, excel_path,
            reasoning_effort=config.ppt_reasoning_effort,
            context_tier=config.ppt_context_tier,
        )
        try:
            if page_number == 2:
                _validate_dashboard_payloads(excel_path, payload, {"sections": []})
            else:
                _validate_dashboard_payloads(
                    excel_path, {"highlights": [], "group_updates": [], "milestones": []}, payload
                )
            return payload
        except ValueError as error:
            if attempt == max_attempts:
                raise ValueError(
                    f"{error} (slide {page_number}, after {max_attempts} attempts)"
                ) from error
            request_prompt = (
                f"{prompt}\n\nSlide {page_number} failed validation: {error}\n"
                "Regenerate the complete JSON object for this slide using the original schema. "
                "Read the cited workbook cells again. Copy dates exactly from their individual "
                "source cells; do not add a year, expand ranges, or put qualifiers in date fields. "
                "Put forecast/actual/worst-case labels only in trend_kind. If a date is not "
                "documented, return null with an empty date source and a null trend_kind when "
                "trend_actual is null. Never fabricate evidence or silently remove supported "
                "facts to bypass validation. Return only valid JSON."
            )


def generate_weekly_ppts(
    config: MrcConfig,
    cycle_code: str,
    excel_artifacts: list[MrcArtifact],
) -> list[MrcArtifact]:
    sources = sorted(
        (artifact for artifact in excel_artifacts if artifact.kind == "excel"),
        key=lambda artifact: artifact.file_name.casefold(),
    )
    if not sources:
        raise FileNotFoundError("No weekly Excel workbooks are available for PPT generation")
    output_dir = WORKSPACE_ROOT / cycle_code / "ppt"
    output_dir.mkdir(parents=True, exist_ok=True)
    generated = []
    for source in sources:
        excel_path = WORKSPACE_ROOT / source.relative_path
        template_path = _template_path_for(source.file_name)
        if not template_path.is_file():
            raise FileNotFoundError(f"PPT template does not exist: {template_path}")
        dashboard = template_path == DMR_16CH_TEMPLATE_PATH
        page2_prompt = DASHBOARD_PAGE2_PROMPT if dashboard else PAGE2_PROMPT
        page3_prompt = DASHBOARD_PAGE3_PROMPT if dashboard else PAGE3_PROMPT
        sheet_names = json.dumps(_sheet_names(excel_path), ensure_ascii=False)
        page2_prompt = page2_prompt.format(cycle_code=cycle_code, sheet_names=sheet_names)
        page3_prompt = page3_prompt.format(
            cycle_code=cycle_code, sheet_names=sheet_names, max_chars=PAGE3_MAX_CHAR_COUNT
        )
        if dashboard:
            page2 = _call_dashboard_page(page2_prompt, config, excel_path, 2)
            page3 = _call_dashboard_page(page3_prompt, config, excel_path, 3)
        else:
            page2 = _call_copilot_json(
                page2_prompt, config.ppt_model, excel_path,
                reasoning_effort=config.ppt_reasoning_effort,
                context_tier=config.ppt_context_tier,
            )
            page3 = _call_copilot_json(
                page3_prompt, config.ppt_model, excel_path,
                reasoning_effort=config.ppt_reasoning_effort,
                context_tier=config.ppt_context_tier,
            )
        page2["source_url"] = source.source_url
        page3["source_url"] = source.source_url
        output_path = output_dir / f"{excel_path.stem}_AI_generated.pptx"
        _update_presentation(template_path, output_path, cycle_code, page2, page3)
        generated.append(register_ppt(cycle_code, output_path, source.file_name))
    return generated