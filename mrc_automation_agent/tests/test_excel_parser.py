import unittest
from io import BytesIO

from openpyxl import Workbook

from mrc_automation_agent.excel_parser import parse_workbook
from mrc_automation_agent.models import WorkbookFile


class ExcelParserTests(unittest.TestCase):
    def test_uses_platform_as_group_for_engineering_domains(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.append(["Platform:\nOKS - RS", "Engineering Domain", "Owner", "Status Comments"])
        sheet.append(["Data Center Component Engineering", "Signal Integrity", "Owner One", "On track"])
        sheet.append([None, "Power Integrity", "Owner Two", "Watch"])
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile("DMR-RS WW40.xlsx", "https://example.invalid/DMR-RS.xlsx", content.getvalue()),
            header_search_rows=10,
        )

        self.assertEqual(["Signal Integrity", "Power Integrity"], [record.project_name for record in records])
        self.assertEqual(
            ["Data Center Component Engineering", "Data Center Component Engineering"],
            [record.function_team for record in records],
        )

    def test_accepts_engineering_domain_as_project_header(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.append(["Engineering Domain", "Owner", "Status Comments"])
        sheet.append(["Signal Integrity", "Example Owner", "On track"])
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile("DMR-RS WW40.xlsx", "https://example.invalid/DMR-RS.xlsx", content.getvalue()),
            header_search_rows=10,
        )

        self.assertEqual(1, len(records))
        self.assertEqual("Signal Integrity", records[0].project_name)

    def test_accepts_status_comments_header_with_instructions(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.append([None, "Platform:\nOKS-AP", "uPLC: Mid Design Entry", "Owner", "Status Comments\n(use this column for updates)"])
        sheet.append([None, "Platform Health", "CPU Package Development", "Example Owner", "On track"])
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile("DMR Dashboard WW38'26.xlsx", "https://example.invalid/DMR.xlsx", content.getvalue()),
            header_search_rows=10,
        )

        self.assertEqual(1, len(records))
        self.assertEqual("CPU Package Development", records[0].project_name)
        self.assertEqual("Platform Health", records[0].function_team)
        self.assertEqual("On track", records[0].status_comments)

    def test_uses_current_week_status_instead_of_history(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.append([None, None, None, None])
        sheet.append([None, None, "WW40'26", "WW39'26"])
        sheet.append(["Engineering Domain", "Owner", "Status Comments", "Status Comments"])
        sheet.append(["Signal Integrity", "owner@example.com", "", "Historical update"])
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile("DMR-RS WW40.xlsx", "https://example.invalid/DMR-RS.xlsx", content.getvalue()),
            header_search_rows=10,
        )

        self.assertEqual("", records[0].status_comments)
        self.assertTrue(records[0].is_missing_update)

    def test_finds_status_comments_by_header_in_column_l(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.title = "Weekly Status"
        sheet.append(["MRC weekly report"])
        headers = [None] * 12
        headers[1] = "Function Team"
        headers[3] = "Project Name"
        headers[6] = "Owner"
        headers[8] = "Owner Email"
        headers[11] = "Status Comments"
        sheet.append(headers)
        values = [None] * 12
        values[1] = "Core Platform"
        values[3] = "Atlas Migration"
        values[6] = "Alex Kim"
        values[8] = "Alex.Kim@example.com"
        values[11] = "On track"
        sheet.append(values)
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile("MRC.xlsx", "https://example.invalid/MRC.xlsx", content.getvalue()),
            header_search_rows=5,
        )

        self.assertEqual(1, len(records))
        self.assertEqual("Weekly Status", records[0].sheet_name)
        self.assertEqual("alex.kim@example.com", records[0].owner_email)
        self.assertEqual("On track", records[0].status_comments)

    def test_reads_name_from_actual_owner_and_platform_headers(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.append([None, "Platform:CRI", None, "Owner", None, None, None, None, None, None, "Status Comments"])
        sheet.append([None, "Project A", None, "Example Owner", None, None, None, None, None, None, ""])
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile("MRC.xlsx", "https://example.invalid/MRC.xlsx", content.getvalue()),
            header_search_rows=5,
        )

        self.assertEqual(1, len(records))
        self.assertEqual("Project A", records[0].project_name)
        self.assertEqual("", records[0].function_team)
        self.assertEqual("Example Owner", records[0].owner_name)
        self.assertEqual("", records[0].owner_email)

    def test_promotes_email_from_owner_column(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.append(["Project: Example", "Owner", "Status Comments"])
        sheet.append(["Project A", "owner@example.com", ""])
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile("MRC.xlsx", "https://example.invalid/MRC.xlsx", content.getvalue()),
            header_search_rows=5,
        )

        self.assertEqual("", records[0].owner_name)
        self.assertEqual("owner@example.com", records[0].owner_email)

    def test_expands_multiple_emails_from_owner_column(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.append(["Project: Example", "Owner", "Status Comments"])
        sheet.append(
            [
                "Project A",
                "first@example.com; second@example.com / third@example.com",
                "",
            ]
        )
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile("MRC.xlsx", "https://example.invalid/MRC.xlsx", content.getvalue()),
            header_search_rows=5,
        )

        self.assertEqual(
            ["first@example.com", "second@example.com", "third@example.com"],
            [record.owner_email for record in records],
        )
        self.assertTrue(all(record.owner_name == "" for record in records))

    def test_uses_target_week_column_for_cor_workbook(self) -> None:
        excel = Workbook()
        sheet = excel.active
        sheet.append(["Project: COR SP", "Owner", "WW24'26", "update for ww24?"])
        sheet.append(["Project A", "Example Owner", "Current status", "Yes"])
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile(
                "COR-SP - DHE Dashboard WW24'26.xlsx",
                "https://example.invalid/COR-SP.xlsx",
                content.getvalue(),
            ),
            header_search_rows=5,
        )

        self.assertEqual(1, len(records))
        self.assertEqual("Example Owner", records[0].owner_name)
        self.assertEqual("Current status", records[0].status_comments)

    def test_finds_week_column_in_underscore_file_name_and_later_sheet(self) -> None:
        excel = Workbook()
        excel.active.title = "Instructions"
        sheet = excel.create_sheet("Execution Status")
        sheet.append(["Project", "Owner Email", "WW15"])
        sheet.append(["Project A", "owner@example.com", "Current status"])
        content = BytesIO()
        excel.save(content)

        records = parse_workbook(
            WorkbookFile(
                "Execution_MRC_WW15.xlsx",
                "https://example.invalid/Execution_MRC_WW15.xlsx",
                content.getvalue(),
            ),
            header_search_rows=5,
        )

        self.assertEqual(1, len(records))
        self.assertEqual("Execution Status", records[0].sheet_name)
        self.assertEqual("Current status", records[0].status_comments)


if __name__ == "__main__":
    unittest.main()