import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import AsyncMock, patch

from openpyxl import Workbook
from pptx import Presentation

from mrc_automation_agent import artifact_store, ppt_generator
from mrc_automation_agent.config import ConfigurationError, MrcConfig
from mrc_automation_agent.models import MrcArtifact


class PptGeneratorTests(unittest.TestCase):
    def test_mrc_ppt_settings_default_when_missing_or_empty(self) -> None:
        for environment in ({}, {"MRC_PPT_REASONING_EFFORT": "", "MRC_PPT_CONTEXT_TIER": " "}):
            with self.subTest(environment=environment), patch.dict("os.environ", environment, clear=True):
                config = MrcConfig.from_env()
                self.assertEqual("high", config.ppt_reasoning_effort)
                self.assertEqual("default", config.ppt_context_tier)

    def test_mrc_ppt_settings_load_and_validate_environment(self) -> None:
        with patch.dict("os.environ", {
            "MRC_PPT_REASONING_EFFORT": " MEDIUM ",
            "MRC_PPT_CONTEXT_TIER": " LONG_CONTEXT ",
        }, clear=True):
            config = MrcConfig.from_env()
            self.assertEqual("medium", config.ppt_reasoning_effort)
            self.assertEqual("long_context", config.ppt_context_tier)
        for name in ("MRC_PPT_REASONING_EFFORT", "MRC_PPT_CONTEXT_TIER"):
            with self.subTest(name=name), patch.dict("os.environ", {name: "invalid"}, clear=True):
                with self.assertRaisesRegex(ConfigurationError, name):
                    MrcConfig.from_env()

    def _dashboard_workbook(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Status WW38"
        sheet["L10"] = "Alpha completed 1211/1211. DC-SCM needs 200 units; requested WW43, forecast WW44."
        sheet["L17"] = "VR testing is expected by WW48 as a worst-case scenario."
        sheet.merge_cells("B5:B8")
        sheet["B5"] = "Engineering"
        workbook.save(path)
        workbook.close()

    def _dashboard_payloads(self) -> tuple[dict, dict]:
        paragraph = {"bold_lead": "DC-SCM.", "normal_detail": "200 units pending.", "source": "Status WW38!L10"}
        return {
            "highlights": [{"bold_lead": "Alpha complete.", "normal_detail": "1211/1211.", "source": "Status WW38!L10"}],
            "group_updates": [{"title": "Manufacturing", "paragraphs": [dict(paragraph)]}],
            "milestones": [{
                "name": "DC-SCM", "drive_to": "WW43", "drive_to_source": "Status WW38!L10",
                "trend_actual": "WW44", "trend_kind": "forecast", "trend_source": "Status WW38!L10",
                "source": "Status WW38!L10",
            }, {
                "name": "VR testing", "drive_to": None, "drive_to_source": "",
                "trend_actual": "WW48", "trend_kind": "worst_case", "trend_source": "Status WW38!L17",
                "source": "Status WW38!L17",
            }],
        }, {"sections": [{"title": "Manufacturing", "paragraphs": [dict(paragraph)]}]}

    def test_selects_template_from_excel_file_name(self) -> None:
        self.assertEqual(
            ppt_generator.DMR_RS_TEMPLATE_PATH,
            ppt_generator._template_path_for("DMR-RS Dashboard.xlsx"),
        )
        self.assertEqual(
            ppt_generator.DMR_RS_TEMPLATE_PATH,
            ppt_generator._template_path_for("dmr-rs dashboard.XLSX"),
        )
        self.assertEqual(
            ppt_generator.DMR_16CH_TEMPLATE_PATH,
            ppt_generator._template_path_for("DMR-16ch Status Dashboard WW43'26.xlsx"),
        )
        self.assertEqual(
            ppt_generator.DMR_16CH_TEMPLATE_PATH,
            ppt_generator._template_path_for("dmr-16CH dashboard.XLSX"),
        )
        self.assertEqual(
            ppt_generator.DEFAULT_TEMPLATE_PATH,
            ppt_generator._template_path_for("Other Dashboard.xlsx"),
        )

    def test_replaces_week_token_with_right_curly_apostrophe_year(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(0, 0, 1000000, 1000000)
        shape.text = "OKS-AP\nWW25’26"

        ppt_generator._replace_week_tokens(slide, "WW40'2026")

        self.assertEqual("OKS-AP\nWW40'2026", shape.text)

    def test_prompts_preserve_platform_and_engineering_domain_hierarchy(self) -> None:
        self.assertIn("Platform is the engineering group", ppt_generator.PAGE2_PROMPT)
        self.assertIn("Engineering Domain is the project", ppt_generator.PAGE2_PROMPT)
        self.assertIn("Organize the narrative by Platform", ppt_generator.PAGE3_PROMPT)

    def test_uses_dmr_16ch_template_preserving_cover_title(self) -> None:
        template_path = ppt_generator._template_path_for(
            "DMR-16ch Status Dashboard WW43'26.xlsx"
        )
        template = Presentation(str(template_path))
        title = next(
            shape for shape in ppt_generator._text_shapes(template.slides[0])
            if "DMR-16ch" in shape.text
        )
        expected_font_size = ppt_generator._first_font_size(title)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output.pptx"
            ppt_generator._update_presentation(
                template_path, output, "2026WW43", {}, {}
            )
            presentation = Presentation(str(output))
            shapes = ppt_generator._text_shapes(presentation.slides[0])
            cover_title = next(shape for shape in shapes if "DMR-16ch" in shape.text)
            cover_text = "\n".join(shape.text for shape in shapes)
            self.assertEqual(expected_font_size, ppt_generator._first_font_size(cover_title))
            self.assertIn("WW43'2026", cover_text)
            self.assertIn("DHE Product Execution Forum", cover_text)
            self.assertNotIn("OKS-AP", cover_text)

    def test_dashboard_fills_named_regions_and_source_links(self) -> None:
        presentation = Presentation(str(ppt_generator.DMR_16CH_TEMPLATE_PATH))
        slide = presentation.slides[1]
        original_bounds = {
            shape.name: (shape.left, shape.top, shape.width, shape.height)
            for shape in slide.shapes
        }
        source_url = "https://example.invalid/status.xlsx?web=1"
        source = "Status WW38!L10"
        ppt_generator._update_page2(slide, "2026WW43", {
            "highlights": [{"bold_lead": "Alpha complete.", "normal_detail": "1211/1211.", "source": source}],
            "group_updates": [{"title": "Manufacturing", "paragraphs": [
                {"bold_lead": "DC-SCM.", "normal_detail": "200 units required.", "source": source}
            ]}],
            "milestones": [{"name": "DC-SCM", "drive_to": "WW43", "trend_actual": "Forecast WW44", "source": source}],
            "source_url": source_url,
        })
        shapes = {shape.name: shape for shape in slide.shapes}
        self.assertIn("Alpha complete.", shapes["highlights"].text)
        self.assertIn("Manufacturing", shapes["group_updates"].text)
        self.assertIn("WW43'2026", shapes["dashboard_title"].text)
        table = shapes["milestone_table"].table
        self.assertEqual("WW43", table.cell(1, 1).text)
        self.assertEqual("Forecast WW44", table.cell(1, 2).text)
        self.assertEqual("243238", str(table.cell(1, 2).text_frame.paragraphs[0].runs[0].font.color.rgb))
        self.assertEqual("", table.cell(2, 0).text)
        for name in ("highlights", "group_updates"):
            links = [run for paragraph in shapes[name].text_frame.paragraphs
                     for run in paragraph.runs if run.hyperlink.address]
            self.assertEqual([" Details"], [run.text for run in links])
            self.assertIn("activeCell=Status%20WW38%21L10", links[0].hyperlink.address)
        self.assertEqual(source_url, shapes["source_workbook_link"].text_frame.paragraphs[0].runs[0].hyperlink.address)
        self.assertEqual(original_bounds, {
            shape.name: (shape.left, shape.top, shape.width, shape.height)
            for shape in slide.shapes
        })

    def test_dashboard_clears_placeholders_without_data_or_url(self) -> None:
        presentation = Presentation(str(ppt_generator.DMR_16CH_TEMPLATE_PATH))
        slide = presentation.slides[1]
        ppt_generator._update_page2(slide, "2026WW43", {})
        shapes = {shape.name: shape for shape in slide.shapes}
        self.assertEqual("No current-week update provided.", shapes["highlights"].text)
        self.assertEqual("No sourced milestones", shapes["milestone_table"].table.cell(1, 0).text)
        self.assertEqual("Source Excel unavailable", shapes["source_workbook_link"].text)
        self.assertNotIn("[", shapes["group_updates"].text)

    def test_date_evidence_allows_only_equivalent_formatting(self) -> None:
        for date, source in (
            ("Dec/January'27", "DVT trending Dec/January\u201927."),
            ("May / June '27", "Fab3 May/June\u201927."),
            ("WW48", "Ready in WW 48 as a worst-case scenario."),
            ("WW42-44", "Expected WW42\u201344."),
        ):
            with self.subTest(date=date):
                self.assertTrue(ppt_generator._date_in_source(date, source))
        for date, source in (
            ("WW4", "Forecast WW44."),
            ("WW45", "Forecast WW44."),
            ("WW44'26", "Forecast WW44."),
            ("January'27", "No milestone date provided."),
            ("Dec/January'27", "DVT trending Dec/January'28."),
        ):
            with self.subTest(date=date, source=source):
                self.assertFalse(ppt_generator._date_in_source(date, source))

    def test_dashboard_date_error_identifies_milestone_value_and_cell(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            excel_path = Path(directory) / "status.xlsx"
            self._dashboard_workbook(excel_path)
            page2, page3 = self._dashboard_payloads()
            page2["milestones"][0]["trend_actual"] = "WW45"
            with self.assertRaises(ValueError) as raised:
                ppt_generator._validate_dashboard_payloads(excel_path, page2, page3)
            for text in ("DC-SCM", "trend_actual", "WW45", "Status WW38!L10"):
                self.assertIn(text, str(raised.exception))

    def test_dashboard_validation_rejects_invalid_sources_and_unfounded_dates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            excel_path = Path(directory) / "status.xlsx"
            self._dashboard_workbook(excel_path)
            for source in ("Missing!L10", "Status WW38!M10", "Status WW38!B6", "Status WW38!L0", "Status WW38!L10:L11"):
                with self.subTest(source=source):
                    page2, page3 = self._dashboard_payloads()
                    page2["highlights"][0]["source"] = source
                    with self.assertRaisesRegex(ValueError, "Invalid DMR-16ch report"):
                        ppt_generator._validate_dashboard_payloads(excel_path, page2, page3)
            for field, value in (("drive_to", "WW45"), ("trend_actual", "WW45"), ("trend_kind", None), ("trend_kind", [])):
                with self.subTest(field=field):
                    page2, page3 = self._dashboard_payloads()
                    page2["milestones"][0][field] = value
                    with self.assertRaises(ValueError):
                        ppt_generator._validate_dashboard_payloads(excel_path, page2, page3)

    def test_dashboard_validation_enforces_capacity_and_allows_missing_dates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            excel_path = Path(directory) / "status.xlsx"
            self._dashboard_workbook(excel_path)
            page2, page3 = self._dashboard_payloads()
            ppt_generator._validate_dashboard_payloads(excel_path, page2, page3)
            page2["milestones"][0].update(trend_actual=None, trend_kind=None, trend_source="")
            ppt_generator._validate_dashboard_payloads(excel_path, page2, page3)
            invalid = deepcopy(page2)
            invalid["highlights"][0]["normal_detail"] = "word " * 100
            with self.assertRaisesRegex(ValueError, "180 characters"):
                ppt_generator._validate_dashboard_payloads(excel_path, invalid, page3)
            invalid = deepcopy(page2)
            invalid["group_updates"][0]["paragraphs"] *= 5
            with self.assertRaisesRegex(ValueError, "at most 4"):
                ppt_generator._validate_dashboard_payloads(excel_path, invalid, page3)

    def test_dashboard_retries_validation_errors_for_only_the_failed_page(self) -> None:
        config = MrcConfig(
            sharepoint_site_url="", sharepoint_folder_template="", sharepoint_tenant="",
            sharepoint_client_id="", sharepoint_thumbprint="", sharepoint_certificate_path="",
            sharepoint_certificate_password="",
        )
        with tempfile.TemporaryDirectory() as directory:
            excel_path = Path(directory) / "status.xlsx"
            self._dashboard_workbook(excel_path)
            page2, page3 = self._dashboard_payloads()
            invalid = deepcopy(page2)
            invalid["milestones"][0]["trend_actual"] = "WW45"
            with patch.object(ppt_generator, "_call_copilot_json", side_effect=[invalid, page2]) as copilot:
                result = ppt_generator._call_dashboard_page("Original slide 2 requirements", config, excel_path, 2)
            self.assertEqual(page2, result)
            self.assertEqual(2, copilot.call_count)
            retry_prompt = copilot.call_args_list[1].args[0]
            for text in ("Original slide 2 requirements", "DC-SCM", "WW45", "Status WW38!L10"):
                self.assertIn(text, retry_prompt)
            invalid_detail = deepcopy(page3)
            invalid_detail["sections"][0]["paragraphs"][0]["source"] = "Status WW38!M10"
            with patch.object(ppt_generator, "_call_copilot_json", side_effect=[invalid_detail, page3]) as copilot:
                result = ppt_generator._call_dashboard_page("Original slide 3 requirements", config, excel_path, 3)
            self.assertEqual(page3, result)
            self.assertEqual(2, copilot.call_count)

    def test_dashboard_stops_after_three_failed_validations(self) -> None:
        config = MrcConfig(
            sharepoint_site_url="", sharepoint_folder_template="", sharepoint_tenant="",
            sharepoint_client_id="", sharepoint_thumbprint="", sharepoint_certificate_path="",
            sharepoint_certificate_password="",
        )
        with tempfile.TemporaryDirectory() as directory:
            excel_path = Path(directory) / "status.xlsx"
            self._dashboard_workbook(excel_path)
            page2, _ = self._dashboard_payloads()
            page2["milestones"][0]["trend_actual"] = "WW45"
            with patch.object(ppt_generator, "_call_copilot_json", return_value=page2) as copilot:
                with self.assertRaisesRegex(ValueError, "slide 2, after 3 attempts"):
                    ppt_generator._call_dashboard_page("Slide 2", config, excel_path, 2)
            self.assertEqual(3, copilot.call_count)
            with patch.object(ppt_generator, "_call_copilot_json", side_effect=RuntimeError("SDK error")) as copilot:
                with self.assertRaisesRegex(RuntimeError, "SDK error"):
                    ppt_generator._call_dashboard_page("Slide 2", config, excel_path, 2)
            self.assertEqual(1, copilot.call_count)

    def test_generates_dashboard_with_dedicated_prompts_and_verified_links(self) -> None:
        config = MrcConfig(
            sharepoint_site_url="", sharepoint_folder_template="", sharepoint_tenant="",
            sharepoint_client_id="", sharepoint_thumbprint="", sharepoint_certificate_path="",
            sharepoint_certificate_password="",
            ppt_reasoning_effort="medium", ppt_context_tier="long_context",
        )
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            excel_dir = workspace / "2026WW43" / "excel"
            excel_dir.mkdir(parents=True)
            excel_path = excel_dir / "DMR-16ch Status Dashboard WW43'26.xlsx"
            self._dashboard_workbook(excel_path)
            artifact = MrcArtifact(
                "excel", excel_path.name, excel_path.relative_to(workspace).as_posix(),
                source_url="https://example.invalid/status.xlsx?web=1",
            )
            page2, page3 = self._dashboard_payloads()
            with (
                patch.object(ppt_generator, "WORKSPACE_ROOT", workspace),
                patch.object(artifact_store, "WORKSPACE_ROOT", workspace),
                patch.object(ppt_generator, "_call_copilot_json", side_effect=[page2, page3]) as copilot,
            ):
                generated = ppt_generator.generate_weekly_ppts(config, "2026WW43", [artifact])
            self.assertIn("one-page MRC executive pre-read", copilot.call_args_list[0].args[0])
            self.assertIn("Status WW38", copilot.call_args_list[0].args[0])
            self.assertIn("Do not pad groups", copilot.call_args_list[1].args[0])
            for call in copilot.call_args_list:
                self.assertEqual({"reasoning_effort": "medium", "context_tier": "long_context"}, call.kwargs)
            deck = Presentation(workspace / generated[0].relative_path)
            shapes = {shape.name: shape for shape in deck.slides[1].shapes}
            table = shapes["milestone_table"].table
            self.assertEqual("Forecast: WW44", table.cell(1, 2).text)
            self.assertEqual("Not provided", table.cell(2, 1).text)
            self.assertEqual("Worst case: WW48", table.cell(2, 2).text)
            self.assertIn("activeCell=Status%20WW38%21L10", table.cell(1, 1).text_frame.paragraphs[0].runs[0].hyperlink.address)
            self.assertEqual(artifact.source_url, shapes["source_workbook_link"].text_frame.paragraphs[0].runs[0].hyperlink.address)
            detail = next(shape for shape in deck.slides[2].shapes if shape.name == "details_body")
            self.assertNotIn("No update available", detail.text)
            self.assertEqual(2, len(detail.text_frame.paragraphs))
            for slide in deck.slides:
                for shape in ppt_generator._text_shapes(slide):
                    self.assertNotIn("[Engineering", shape.text)
                    self.assertNotIn("link pending", shape.text)

    def test_copilot_call_does_not_inherit_log_analysis_workspace(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            excel_path = Path(directory) / "status.xlsx"
            excel_path.write_bytes(b"test")
            copilot = AsyncMock(return_value='{"executive_summary": []}')

            with patch.object(ppt_generator, "ask_copilot", copilot):
                ppt_generator._call_copilot_json("Analyze workbook", "gpt-5.5", excel_path)

            copilot.assert_awaited_once_with(
                "Analyze workbook",
                model="gpt-5.5",
                files=[excel_path],
                workspace="",
                reasoning_effort="high",
                context_tier="default",
            )

    def test_generates_one_presentation_for_each_workbook(self) -> None:
        config = MrcConfig(
            sharepoint_site_url="",
            sharepoint_folder_template="",
            sharepoint_tenant="",
            sharepoint_client_id="",
            sharepoint_thumbprint="",
            sharepoint_certificate_path="",
            sharepoint_certificate_password="",
        )
        page2 = {
            "executive_summary": [
                {"bold_lead": "On track.", "normal_detail": "Most work is green."},
                {"bold_lead": "One risk.", "normal_detail": "Material is late."},
            ],
            "detail_sections": [],
        }
        page3 = {
            "sections": [
                {
                    "title": "Engineering",
                    "paragraphs": [
                        {
                            "bold_lead": "Stable.",
                            "normal_detail": "Milestones met.",
                            "source": "Status WW38!L12",
                        },
                        {
                            "bold_lead": "Watch.",
                            "normal_detail": "Supplier timing.",
                            "source": "Status WW38!L18",
                        },
                    ],
                }
            ]
        }
        template = Presentation(str(ppt_generator.DEFAULT_TEMPLATE_PATH))
        template_progress_labels = [
            shape.text
            for shape in template.slides[1].shapes
            if getattr(shape, "has_text_frame", False)
            and shape.text.startswith("WW")
        ]

        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            excel_dir = workspace / "2026WW38" / "excel"
            excel_dir.mkdir(parents=True)
            execution = excel_dir / "CRI-ww38-MRC.xlsx"
            tracking = excel_dir / "DHE Execution MRC AR Tracking.xlsx"
            for path in (execution, tracking):
                path.write_bytes(b"test")
            artifacts = [
                MrcArtifact(
                    "excel",
                    execution.name,
                    execution.relative_to(workspace).as_posix(),
                    source_url="https://example.invalid/execution.xlsx?web=1",
                ),
                MrcArtifact(
                    "excel",
                    tracking.name,
                    tracking.relative_to(workspace).as_posix(),
                    source_url="https://example.invalid/tracking.xlsx?web=1",
                ),
            ]
            with (
                patch.object(ppt_generator, "WORKSPACE_ROOT", workspace),
                patch.object(artifact_store, "WORKSPACE_ROOT", workspace),
                patch.object(
                    ppt_generator,
                    "_sheet_names",
                    side_effect=[["Status WW38", "Lookup"], ["AR tracking"]],
                ),
                patch.object(
                    ppt_generator,
                    "_call_copilot_json",
                    side_effect=[page2, page3, page2, page3],
                ) as copilot,
            ):
                generated = ppt_generator.generate_weekly_ppts(
                    config, "2026WW38", artifacts
                )

            self.assertEqual([execution.name, tracking.name], [item.source_workbook_name for item in generated])
            self.assertEqual(
                [
                    "CRI-ww38-MRC_AI_generated.pptx",
                    "DHE Execution MRC AR Tracking_AI_generated.pptx",
                ],
                [item.file_name for item in generated],
            )
            self.assertEqual(execution, copilot.call_args_list[0].args[2])
            self.assertEqual(tracking, copilot.call_args_list[2].args[2])
            self.assertIn('["Status WW38", "Lookup"]', copilot.call_args_list[0].args[0])
            self.assertIn('["AR tracking"]', copilot.call_args_list[2].args[0])
            expected_source_urls = {
                execution.name: "https://example.invalid/execution.xlsx?web=1",
                tracking.name: "https://example.invalid/tracking.xlsx?web=1",
            }
            for artifact in generated:
                output = workspace / artifact.relative_path
                self.assertTrue(output.is_file())
                presentation = Presentation(output)
                self.assertEqual(3, len(presentation.slides))
                page2_text_shapes = [
                    shape
                    for shape in presentation.slides[1].shapes
                    if getattr(shape, "has_text_frame", False)
                ]
                self.assertIn(
                    "Platform Dashboard –WW38'2026",
                    [shape.text for shape in page2_text_shapes],
                )
                dashboard_shape = next(
                    shape
                    for shape in page2_text_shapes
                    if shape.text.startswith("Platform Dashboard")
                )
                self.assertEqual(
                    32.0,
                    dashboard_shape.text_frame.paragraphs[0].runs[0].font.size.pt,
                )
                self.assertEqual(
                    template_progress_labels,
                    [shape.text for shape in page2_text_shapes if shape.text.startswith("WW")],
                )
                self.assertTrue(
                    all(
                        shape.text_frame.paragraphs[0].runs[0].font.size.pt == 8.0
                        for shape in page2_text_shapes
                        if shape.text.startswith("WW")
                    )
                )
                detail_shapes = [
                    shape
                    for shape in page2_text_shapes
                    if shape.left == ppt_generator.PAGE2_DETAIL_TEXTBOX_LEFT
                    and shape.top == ppt_generator.PAGE2_DETAIL_TEXTBOX_TOP
                ]
                self.assertEqual(1, len(detail_shapes))
                hyperlinks = [
                    run.hyperlink.address
                    for shape in presentation.slides[2].shapes
                    if getattr(shape, "has_text_frame", False)
                    for paragraph in shape.text_frame.paragraphs
                    for run in paragraph.runs
                    if run.hyperlink.address
                ]
                self.assertEqual(2, len(hyperlinks))
                self.assertTrue(
                    hyperlinks[0].startswith(expected_source_urls[artifact.source_workbook_name])
                )
                self.assertIn("activeCell=Status%20WW38%21L12", hyperlinks[0])
                self.assertIn("wdActiveCell=Status%20WW38%21L12", hyperlinks[0])

    def test_source_link_opens_sharepoint_workbook_in_web_mode(self) -> None:
        link = ppt_generator._source_link(
            "https://example.invalid/MRC%20WW38.xlsx", "Status WW38!L12"
        )

        self.assertEqual(
            "https://example.invalid/MRC%20WW38.xlsx?web=1"
            "&activeCell=Status%20WW38%21L12"
            "&wdActiveCell=Status%20WW38%21L12",
            link,
        )


if __name__ == "__main__":
    unittest.main()