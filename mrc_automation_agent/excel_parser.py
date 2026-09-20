import re
from io import BytesIO
from typing import Any

from .models import ProjectRecord, WorkbookFile


class WorkbookFormatError(ValueError):
    pass


HEADER_ALIASES = {
    "function_team": {"functionteam", "group", "team"},
    "project_name": {"engineeringdomain", "project", "projectname"},
    "owner_name": {"owner", "ownername", "projectowner"},
    "owner_email": {"owneremail", "projectowneremail", "email"},
    "status_comments": {"comments", "statuscomment", "statuscomments"},
}
REQUIRED_HEADERS = {"project_name", "status_comments"}
EMAIL_PATTERN = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
EMAIL_SEPARATOR_PATTERN = re.compile(r"[\s,;/]+")


def _email_addresses(value: str) -> list[str]:
    parts = [part.lower() for part in EMAIL_SEPARATOR_PATTERN.split(value.strip()) if part]
    return parts if parts and all(EMAIL_PATTERN.fullmatch(part) for part in parts) else []


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _header_name(value: Any) -> str | None:
    text = str(value or "").strip().lower()
    if re.match(r"^platform\s*:", text):
        return "platform"
    if re.match(r"^(project|uplc)\s*:", text):
        return "project_name"
    if re.match(r"^status\s*comments?\b", text):
        return "status_comments"
    normalized = _normalize_header(value)
    return next(
        (name for name, aliases in HEADER_ALIASES.items() if normalized in aliases),
        None,
    )


def _target_week(workbook_name: str) -> str | None:
    match = re.search(r"WW\s*0?([1-9]|[1-4]\d|5[0-3])\b", workbook_name, re.IGNORECASE)
    return match.group(1) if match else None


def _find_headers(
    worksheet, search_rows: int, workbook_name: str
) -> tuple[int, dict[str, int]]:
    target_week = _target_week(workbook_name)
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=search_rows, values_only=True),
        start=1,
    ):
        columns: dict[str, int] = {}
        for column, value in enumerate(row):
            name = _header_name(value)
            if name is not None:
                if name == "status_comments":
                    parent_headers = (
                        worksheet.cell(row=parent_row, column=column + 1).value
                        for parent_row in range(1, row_number)
                    )
                    is_target_week = target_week and any(
                        re.match(
                            rf"^WW\s*0?{re.escape(target_week)}(?:\D|$)",
                            str(parent_header or "").strip(),
                            re.IGNORECASE,
                        )
                        for parent_header in parent_headers
                    )
                    if is_target_week or "status_comments" not in columns:
                        columns[name] = column
                else:
                    columns[name] = column
                continue
            text = str(value or "").strip()
            if target_week and re.match(
                rf"^WW\s*0?{re.escape(target_week)}(?:\D|$)", text, re.IGNORECASE
            ):
                columns["status_comments"] = column
        platform_column = columns.get("platform")
        if platform_column is not None:
            if "project_name" in columns:
                columns.setdefault("function_team", platform_column)
            else:
                columns["project_name"] = platform_column
        has_owner = "owner_name" in columns or "owner_email" in columns
        if REQUIRED_HEADERS.issubset(columns) and has_owner:
            return row_number, columns
    raise WorkbookFormatError(
        f"Could not find required headers in the first {search_rows} rows"
    )


def parse_workbook(workbook: WorkbookFile, header_search_rows: int) -> list[ProjectRecord]:
    from openpyxl import load_workbook

    excel = load_workbook(BytesIO(workbook.content), read_only=True, data_only=True)
    try:
        if not excel.worksheets:
            raise WorkbookFormatError(f"Workbook has no worksheets: {workbook.name}")
        records: list[ProjectRecord] = []
        parsed_sheet = False
        for worksheet in excel.worksheets:
            try:
                header_row, columns = _find_headers(
                    worksheet, header_search_rows, workbook.name
                )
            except WorkbookFormatError:
                continue
            parsed_sheet = True
            current_function_team = ""
            for source_row, values in enumerate(
                worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                def value(name: str) -> str:
                    column = columns.get(name)
                    if column is None or column >= len(values):
                        return ""
                    return str(values[column] or "").strip()

                project_name = value("project_name")
                if not project_name:
                    continue
                function_team = value("function_team")
                if function_team:
                    current_function_team = function_team
                owner_name = value("owner_name")
                owner_emails = _email_addresses(value("owner_email"))
                if not owner_emails:
                    owner_emails = _email_addresses(owner_name)
                if owner_emails and not value("owner_email"):
                    owner_name = ""
                for owner_email in owner_emails or [""]:
                    records.append(
                        ProjectRecord(
                            workbook_name=workbook.name,
                            workbook_url=workbook.source_url,
                            sheet_name=worksheet.title,
                            source_row=source_row,
                            function_team=current_function_team,
                            project_name=project_name,
                            owner_name=owner_name,
                            owner_email=owner_email,
                            status_comments=value("status_comments"),
                        )
                    )
        if not parsed_sheet:
            raise WorkbookFormatError(
                f"Could not find a supported worksheet in {workbook.name}"
            )
        return records
    finally:
        excel.close()
